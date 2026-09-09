import os
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from .common import (
    add_chapter_title, add_heading_1, add_heading_2, add_heading_3,
    add_body_paragraph, add_code_block, set_cell_margins, set_cell_shading,
    find_source_code_file
)

def build_appendix(doc):
    doc.add_page_break()
    add_chapter_title(doc, "Appendix", "Source Code Listing and Technical Specifications")

    add_body_paragraph(doc, "This appendix provides an exhaustive technical specification, structural breakdown, and complete authentic source code listing for all six components of the VetiCare Veterinary Clinic Management System codebase. Every section details the operational purpose, architectural role, entity attributes, UI element mappings, functional algorithms, and execution logic prior to presenting the un-truncated source code.")

    # =========================================================================
    # A.1 Database Schema DDL Specification
    # =========================================================================
    add_heading_1(doc, "A.1 Relational Database Schema Specification & DDL (`veticare.db`)")
    
    add_heading_2(doc, "A.1.1 Architectural Overview & RDBMS Engine Configuration")
    add_body_paragraph(doc, "The database tier of VetiCare is implemented as an embedded relational SQLite database stored in a single cross-platform file named `veticare.db`. Operating directly within the application process space, SQLite eliminates client-server network socket overhead and daemon configuration while providing complete Atomicity, Consistency, Isolation, and Durability (ACID) compliance.")
    add_body_paragraph(doc, "To enforce structural data integrity across clinical operations, referential integrity is explicitly enabled upon opening every database connection via the SQLite pragma command: `PRAGMA foreign_keys = ON;`. Write operations utilize Write-Ahead Logging (WAL) journal mode, permitting concurrent read queries from multiple client browser tabs without encountering database table lock contention.")

    add_heading_2(doc, "A.1.2 Relational Entity & Table Specifications")
    add_body_paragraph(doc, "The database schema consists of five normalized tables structured in Third Normal Form (3NF):")

    add_body_paragraph(doc, "1. `owners` Table: Serves as the primary client directory, storing demographic and contact details for pet owners. Each owner record is identified by a unique text primary key (`id` formatted as `o-xxxx`). The table maintains a one-to-many relationship with registered pet patients.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Table A.1.1: Detailed Column Specification for `owners` Table")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.font.bold = True

    t_o = doc.add_table(rows=1, cols=6)
    t_o.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_o = t_o.rows[0].cells
    h_o[0].text = "Column"
    h_o[1].text = "Data Type"
    h_o[2].text = "Nullability"
    h_o[3].text = "Key Constraint"
    h_o[4].text = "Functional Purpose"
    h_o[5].text = "Sample Value"
    for c in h_o:
        set_cell_shading(c, "E6E6E6")
        c.paragraphs[0].runs[0].font.bold = True

    owners_cols = [
        ("id", "TEXT", "NOT NULL", "PRIMARY KEY", "Unique owner identifier string", "'o-1'"),
        ("name", "TEXT", "NOT NULL", "None", "Full legal name of client", "'Eleanor Vance'"),
        ("phone", "TEXT", "NOT NULL", "None", "Primary contact telephone", "'(555) 234-5678'"),
        ("email", "TEXT", "NOT NULL", "None", "Contact email address", "'eleanor@example.com'"),
        ("address", "TEXT", "NOT NULL", "None", "Physical residential address", "'124 Hill House Way'"),
        ("emergencyContact", "TEXT", "NULLABLE", "None", "Secondary emergency contact", "'Thomas Vance (Brother)'"),
        ("createdAt", "TEXT", "NOT NULL", "None", "Registration ISO timestamp", "'2026-01-15T09:30:00Z'")
    ]
    for col, dt, nl, kc, fp, sv in owners_cols:
        rc = t_o.add_row().cells
        rc[0].text = col
        rc[1].text = dt
        rc[2].text = nl
        rc[3].text = kc
        rc[4].text = fp
        rc[5].text = sv
        rc[0].paragraphs[0].runs[0].font.bold = True
        for i in range(6):
            set_cell_margins(rc[i], top=25, bottom=25, left=45, right=45)
            p_c = rc[i].paragraphs[0]
            p_c.paragraph_format.line_spacing = 1.15
            p_c.paragraph_format.space_after = Pt(2)
            p_c.runs[0].font.size = Pt(8.5)

    add_body_paragraph(doc, "2. `pets` Table: Stores patient profiles linked to an owner via foreign key `ownerId`. Includes species, breed, DOB, weight, medical notes, and vaccination status badge ('Up-to-date' vs 'Not Vaccinated'). Defines declarative cascade deletion (`FOREIGN KEY(ownerId) REFERENCES owners(id) ON DELETE CASCADE`).")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Table A.1.2: Detailed Column Specification for `pets` Table")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.font.bold = True

    t_p = doc.add_table(rows=1, cols=6)
    t_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_p = t_p.rows[0].cells
    h_p[0].text = "Column"
    h_p[1].text = "Data Type"
    h_p[2].text = "Nullability"
    h_p[3].text = "Key Constraint"
    h_p[4].text = "Functional Purpose"
    h_p[5].text = "Sample Value"
    for c in h_p:
        set_cell_shading(c, "E6E6E6")
        c.paragraphs[0].runs[0].font.bold = True

    pets_cols = [
        ("id", "TEXT", "NOT NULL", "PRIMARY KEY", "Unique pet patient identifier", "'p-1'"),
        ("ownerId", "TEXT", "NOT NULL", "FOREIGN KEY (CASCADE)", "Reference to parent owner record", "'o-1'"),
        ("name", "TEXT", "NOT NULL", "None", "Name of pet patient", "'Barnaby'"),
        ("species", "TEXT", "NOT NULL", "None", "Animal species classification", "'Cat'"),
        ("breed", "TEXT", "NOT NULL", "None", "Specific breed classification", "'Domestic Shorthair'"),
        ("dob", "TEXT", "NOT NULL", "None", "Date of birth (YYYY-MM-DD)", "'2021-04-12'"),
        ("weight", "REAL", "NOT NULL", "None", "Body weight in kilograms", "4.8"),
        ("vaccination", "TEXT", "NOT NULL", "None", "Immunization status badge", "'Up-to-date'"),
        ("notes", "TEXT", "NULLABLE", "None", "Clinical alerts or medical notes", "'Mild allergy to penicillin'")
    ]
    for col, dt, nl, kc, fp, sv in pets_cols:
        rc = t_p.add_row().cells
        rc[0].text = col
        rc[1].text = dt
        rc[2].text = nl
        rc[3].text = kc
        rc[4].text = fp
        rc[5].text = sv
        rc[0].paragraphs[0].runs[0].font.bold = True
        for i in range(6):
            set_cell_margins(rc[i], top=25, bottom=25, left=45, right=45)
            p_c = rc[i].paragraphs[0]
            p_c.paragraph_format.line_spacing = 1.15
            p_c.paragraph_format.space_after = Pt(2)
            p_c.runs[0].font.size = Pt(8.5)

    add_body_paragraph(doc, "3. `appointments` Table: Tracks consultation bookings, linking `petId` to attending veterinarian, date/time slot, consultation reason, and status ('Scheduled', 'Completed', 'Cancelled').")
    add_body_paragraph(doc, "4. `treatments` Table: Stores clinical examination records logged by veterinarians, capturing symptoms, diagnosis, prescribed medications, vaccine shot administered, and procedure cost.")
    add_body_paragraph(doc, "5. `invoices` Table: Stores itemized billing transactions generated automatically upon treatment logging. Records consultation fee ($50.00), treatment cost, medication cost, 8% tax, total amount, payment status ('Pending'/'Paid'), and payment method ('Card'/'Cash').")

    add_heading_2(doc, "A.1.3 SQL Data Definition Language (DDL) Script")
    add_body_paragraph(doc, "The authentic SQL DDL script executed during SQLite database initialization is listed below:")

    sql_schema = """-- VetiCare Relational SQLite Database Schema (veticare.db)
PRAGMA foreign_keys = ON;

-- 1. Owners Table
CREATE TABLE IF NOT EXISTS owners (
    id TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    phone TEXT NOT NULL,
    email TEXT NOT NULL,
    address TEXT NOT NULL,
    emergencyContact TEXT,
    createdAt TEXT NOT NULL
);

-- 2. Pets Table (Patient Profiles)
CREATE TABLE IF NOT EXISTS pets (
    id TEXT PRIMARY KEY,
    ownerId TEXT NOT NULL,
    name TEXT NOT NULL,
    species TEXT NOT NULL,
    breed TEXT NOT NULL,
    dob TEXT NOT NULL,
    weight REAL NOT NULL,
    vaccination TEXT NOT NULL,
    notes TEXT,
    FOREIGN KEY(ownerId) REFERENCES owners(id) ON DELETE CASCADE
);

-- 3. Appointments Table (Scheduler)
CREATE TABLE IF NOT EXISTS appointments (
    id TEXT PRIMARY KEY,
    petId TEXT NOT NULL,
    vetName TEXT NOT NULL,
    dateTime TEXT NOT NULL,
    reason TEXT NOT NULL,
    status TEXT NOT NULL,
    FOREIGN KEY(petId) REFERENCES pets(id) ON DELETE CASCADE
);

-- 4. Treatments Table (Clinical Visit Logs)
CREATE TABLE IF NOT EXISTS treatments (
    id TEXT PRIMARY KEY,
    appointmentId TEXT,
    petId TEXT NOT NULL,
    dateTime TEXT NOT NULL,
    symptoms TEXT NOT NULL,
    diagnosis TEXT NOT NULL,
    medications TEXT,
    vaccineGiven TEXT,
    cost REAL NOT NULL,
    notes TEXT,
    FOREIGN KEY(petId) REFERENCES pets(id) ON DELETE CASCADE
);

-- 5. Invoices Table (Financial Transactions)
CREATE TABLE IF NOT EXISTS invoices (
    id TEXT PRIMARY KEY,
    treatmentId TEXT,
    petId TEXT NOT NULL,
    dateTime TEXT NOT NULL,
    consultationFee REAL NOT NULL,
    treatmentCost REAL NOT NULL,
    medicineCost REAL NOT NULL,
    tax REAL NOT NULL,
    total REAL NOT NULL,
    status TEXT NOT NULL,
    paymentMethod TEXT NOT NULL,
    reason TEXT NOT NULL,
    FOREIGN KEY(petId) REFERENCES pets(id) ON DELETE CASCADE
);"""
    add_code_block(doc, sql_schema)

    # =========================================================================
    # A.2 Presentation Layer Layout Specification (index.html)
    # =========================================================================
    doc.add_page_break()
    add_heading_1(doc, "A.2 Presentation Layer Layout & View Specification (`index.html`)")

    add_heading_2(doc, "A.2.1 Architectural Overview & View Panel Breakdown")
    add_body_paragraph(doc, "The presentation layer is implemented as a single HTML5 document (`index.html`) establishing the DOM container layout for the Single Page Application (SPA). Navigational view panels are defined as sibling container sections (`<section id=\"view-xxx\" class=\"view-panel\">`) toggled dynamically by the client router.")

    add_body_paragraph(doc, "Component Layout Summary:\n"
                            "1. Sidebar Navigation (`<aside class=\"sidebar\">`): Contains clinic branding logo, navigation links (`#dashboard`, `#owners`, `#pets`, `#appointments`, `#billing`), and visual theme toggle button (`#theme-toggle`).\n"
                            "2. Dashboard View (`#view-dashboard`): Renders 4 metric summary cards (Total Owners, Total Pets, Today's Appointments, Total Revenue) and recent clinical activity timeline.\n"
                            "3. Owner Directory View (`#view-owners`): Provides search filter input `#owner-search`, '+ Add Owner' modal trigger, and owner card grid container `#owners-grid`.\n"
                            "4. Pet Patient Directory View (`#view-pets`): Features species/breed filter dropdowns, pet card grid container `#pets-grid`, and patient profile details view `#view-pet-details`.\n"
                            "5. Appointment Scheduler View (`#view-appointments`): Renders daily agenda table `#appointments-table`, date picker, vet selector, and status update actions.\n"
                            "6. Billing & Financial Ledger View (`#view-billing`): Displays revenue summary metrics, financial transactions table `#billing-table`, payment action triggers, and receipt modal launcher.\n"
                            "7. Modal Form Dialogs (`#owner-modal`, `#pet-modal`, `#appointment-modal`, `#treatment-modal`): Form overlays capturing registration and clinical visit inputs.\n"
                            "8. Thermal POS Receipt View (`#invoice-modal`): Isolated 300px thermal receipt template formatted for POS browser printing.")

    add_heading_2(doc, "A.2.2 Complete `index.html` Source Code Listing")
    filepath_html = find_source_code_file("index.html")
    if filepath_html and os.path.exists(filepath_html):
        with open(filepath_html, 'r', encoding='utf-8') as f:
            add_code_block(doc, f.read())

    # =========================================================================
    # A.3 Styling System Specification (style.css)
    # =========================================================================
    doc.add_page_break()
    add_heading_1(doc, "A.3 Styling System & Thermal Print Engine Specification (`style.css`)")

    add_heading_2(doc, "A.3.1 Design System Tokens & Thermal Print Rules")
    add_body_paragraph(doc, "The styling tier (`style.css`) provides a complete Vanilla CSS design system built on CSS Custom Properties (variables) for instant Dark/Light theme switching.")

    add_body_paragraph(doc, "Key Styling Modules:\n"
                            "• Design Tokens (`:root`): Defines primary theme colors (`--primary: #2563eb`), card backgrounds (`--bg-card`), text colors (`--text-main`), border radii (`--radius: 12px`), and transition timings.\n"
                            "• Dark Mode Engine (`[data-theme=\"dark\"]`): Overrides background variables (`--bg-main: #0f172a`, `--bg-card: #1e293b`), delivering high visual contrast for dim clinical exam rooms.\n"
                            "• Responsive CSS Grid & Flexbox Rules: Implements dynamic responsive grids (`grid-template-columns: repeat(auto-fill, minmax(280px, 1fr))`) ensuring tablet and desktop compatibility.\n"
                            "• POS Thermal Print Engine (`@media print`): Hides web application controls, sidebars, and headers (`display: none !important`), isolating the 300px monospaced receipt modal `#invoice-modal` for thermal printers.")

    add_heading_2(doc, "A.3.2 Complete `style.css` Source Code Listing")
    filepath_css = find_source_code_file("style.css")
    if filepath_css and os.path.exists(filepath_css):
        with open(filepath_css, 'r', encoding='utf-8') as f:
            add_code_block(doc, f.read())

    # =========================================================================
    # A.4 Single Page Application Engine Specification (app.js)
    # =========================================================================
    doc.add_page_break()
    add_heading_1(doc, "A.4 Single Page Application Client Engine Specification (`app.js`)")

    add_heading_2(doc, "A.4.1 Client Architecture & Function Registry")
    add_body_paragraph(doc, "The client application logic (`app.js`) is organized into a modular namespace object (`window.app`) managing state, HTTP requests, URI hash routing, and DOM view rendering.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Table A.4.1: JavaScript Client Function Registry & Operational Roles")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.font.bold = True

    t_js = doc.add_table(rows=1, cols=4)
    t_js.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_j = t_js.rows[0].cells
    h_j[0].text = "Module / Method"
    h_j[1].text = "Input Parameters"
    h_j[2].text = "Return / State Mutation"
    h_j[3].text = "Functional Responsibility"
    for c in h_j:
        set_cell_shading(c, "E6E6E6")
        c.paragraphs[0].runs[0].font.bold = True

    js_funcs = [
        ("`app.init()`", "None", "Initializes app lifecycle", "Binds hash routing events, sets theme toggle, triggers data load."),
        ("`app.router()`", "None", "Updates active view DOM", "Parses `location.hash`, updates nav active states, renders active panel."),
        ("`app.fetchData()`", "None", "Populates `app.state`", "Executes `GET /api/data`, updates local state arrays, triggers view renders."),
        ("`app.postAPI(url, data)`", "Endpoint URL, Object", "Promise (JSON result)", "Wrapper for asynchronous POST requests with error handling."),
        ("`app.putAPI(url, data)`", "Endpoint URL, Object", "Promise (JSON result)", "Wrapper for asynchronous PUT status update requests."),
        ("`app.owners.render()`", "None", "Renders `#owners-grid`", "Filters owner list by search term, renders card grid HTML."),
        ("`app.pets.render()`", "None", "Renders `#pets-grid`", "Evaluates age from DOB, renders species cards & vaccine badges."),
        ("`app.appointments.render()`", "None", "Renders `#appointments-table`", "Compiles daily agenda table with vet names and status action buttons."),
        ("`app.billing.render()`", "None", "Renders `#billing-table`", "Calculates revenue metrics, renders financial table & payment buttons."),
        ("`app.billing.showReceipt(id)`", "Invoice ID string", "Opens `#invoice-modal`", "Formats itemized fee math and POS thermal receipt layout."),
        ("`app.utils.escapeHTML(str)`", "String", "Sanitized String", "Escapes HTML special characters preventing XSS injection.")
    ]
    for m_m, i_p, r_s, f_r in js_funcs:
        rc = t_js.add_row().cells
        rc[0].text = m_m
        rc[1].text = i_p
        rc[2].text = r_s
        rc[3].text = f_r
        rc[0].paragraphs[0].runs[0].font.bold = True
        for i in range(4):
            set_cell_margins(rc[i], top=25, bottom=25, left=45, right=45)
            p_c = rc[i].paragraphs[0]
            p_c.paragraph_format.line_spacing = 1.15
            p_c.paragraph_format.space_after = Pt(2)
            p_c.runs[0].font.size = Pt(8.5)

    add_heading_2(doc, "A.4.2 Complete `app.js` Source Code Listing")
    filepath_js = find_source_code_file("app.js")
    if filepath_js and os.path.exists(filepath_js):
        with open(filepath_js, 'r', encoding='utf-8') as f:
            add_code_block(doc, f.read())

    # =========================================================================
    # A.5 Python REST API Backend Specification (server.py)
    # =========================================================================
    doc.add_page_break()
    add_heading_1(doc, "A.5 Python REST API Backend Server Specification (`server.py`)")

    add_heading_2(doc, "A.5.1 Server Architecture & Request Routing Logic")
    add_body_paragraph(doc, "The Python backend (`server.py`) implements a zero-dependency REST API server using native standard library modules (`http.server`, `sqlite3`, `json`). Bound to `127.0.0.1:8000`, the server handles request routing, multi-table SQL transactions, auto-invoicing fee calculations, and vaccine status synchronization.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Table A.5.1: Python Backend Endpoint Handler Specifications")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.font.bold = True

    t_py = doc.add_table(rows=1, cols=4)
    t_py.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_py = t_py.rows[0].cells
    h_py[0].text = "HTTP Verb"
    h_py[1].text = "URI Route"
    h_py[2].text = "SQL / Business Action"
    h_py[3].text = "JSON Output Response Payload"
    for c in h_py:
        set_cell_shading(c, "E6E6E6")
        c.paragraphs[0].runs[0].font.bold = True

    py_routes = [
        ("GET", "/api/data", "SELECT * from all 5 tables", "`{ owners:[], pets:[], appointments:[], treatments:[], invoices:[] }`"),
        ("POST", "/api/owners", "INSERT INTO owners ...", "`{ status: 'success', id: 'o-xxxx' }`"),
        ("POST", "/api/pets", "INSERT INTO pets ...", "`{ status: 'success', id: 'p-xxxx' }`"),
        ("POST", "/api/appointments", "INSERT INTO appointments ...", "`{ status: 'success', id: 'a-xxxx' }`"),
        ("POST", "/api/treatments", "BEGIN TRANSACTION; INSERT treatment; INSERT invoice; UPDATE pet vaccine; COMMIT;", "`{ status: 'success', id: 't-xxxx', invoiceId: 'INV-xxxx' }`"),
        ("PUT", "/api/appointments", "UPDATE appointments SET status=?", "`{ status: 'success' }`"),
        ("PUT", "/api/invoices", "UPDATE invoices SET status='Paid'", "`{ status: 'success' }`"),
        ("OPTIONS", "*", "Inject CORS HTTP headers", "HTTP 200 OK empty response with CORS headers")
    ]
    for verb, route, act, resp in py_routes:
        rc = t_py.add_row().cells
        rc[0].text = verb
        rc[1].text = route
        rc[2].text = act
        rc[3].text = resp
        rc[0].paragraphs[0].runs[0].font.bold = True
        for i in range(4):
            set_cell_margins(rc[i], top=25, bottom=25, left=45, right=45)
            p_c = rc[i].paragraphs[0]
            p_c.paragraph_format.line_spacing = 1.15
            p_c.paragraph_format.space_after = Pt(2)
            p_c.runs[0].font.size = Pt(8.5)

    add_heading_2(doc, "A.5.2 Complete `server.py` Source Code Listing")
    filepath_py = find_source_code_file("server.py")
    if filepath_py and os.path.exists(filepath_py):
        with open(filepath_py, 'r', encoding='utf-8') as f:
            add_code_block(doc, f.read())

    # =========================================================================
    # A.6 PHP PDO REST API Backend Specification (api.php)
    # =========================================================================
    doc.add_page_break()
    add_heading_1(doc, "A.6 PHP PDO REST API Backend Server Specification (`api.php`)")

    add_heading_2(doc, "A.6.1 Alternative Runtime Architecture & PDO Features")
    add_body_paragraph(doc, "The PHP backend (`api.php`) provides an alternative zero-dependency API implementation for standard Web hosting environments (Apache/Nginx with PHP). Utilizing PHP Data Objects (PDO), it enforces prepared statement parameter binding (`$stmt->execute([$param])`), avoiding SQL injection risks while providing complete multi-table transaction atomicity.")

    add_heading_2(doc, "A.6.2 Complete `api.php` Source Code Listing")
    filepath_php = find_source_code_file("api.php")
    if filepath_php and os.path.exists(filepath_php):
        with open(filepath_php, 'r', encoding='utf-8') as f:
            add_code_block(doc, f.read())
